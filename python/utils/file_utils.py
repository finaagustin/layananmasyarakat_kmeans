from pathlib import Path

import openpyxl
import pandas as pd


# ==========================================================
# VALIDASI FILE
# ==========================================================

def validate_excel_file(file_path):
    """
    Memastikan file tersedia dan memiliki format Excel.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"File dataset tidak ditemukan: {file_path}"
        )

    if not file_path.is_file():
        raise ValueError(
            f"Path dataset bukan file: {file_path}"
        )

    if file_path.suffix.lower() not in [
        ".xlsx",
        ".xls",
    ]:
        raise ValueError(
            "Dataset harus berformat XLSX atau XLS."
        )

    return file_path


# ==========================================================
# MENENTUKAN JENIS DATASET
# ==========================================================

def determine_source(file_path):
    """
    Menentukan jenis dataset berdasarkan nama file.
    """

    file_name = Path(
        file_path
    ).name.lower()

    if "maret" in file_name:
        return "Maret2026"

    if "april" in file_name:
        return "April2026"

    if "mei" in file_name:
        return "Mei2026"

    if "pemakaman" in file_name:
        return "Pemakaman"

    raise ValueError(
        f"Jenis dataset tidak dikenali: "
        f"{Path(file_path).name}"
    )


# ==========================================================
# MEMBACA DATASET MARET
# ==========================================================

def read_maret(file_path):
    """
    Membaca Data_Gabungan_Maret_2026.xlsx.

    Struktur mengikuti 01_preprocessing.py penelitian:
    - Sheet    : Maret 2026
    - Tanggal  : kolom A
    - Kategori : kolom E
    - Teks     : kolom F
    - Mulai    : baris 2
    """

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    sheet_name = "Maret 2026"

    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' tidak ditemukan "
            f"pada dataset Maret."
        )

    worksheet = workbook[
        sheet_name
    ]

    rows = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        tanggal = row[0]
        kategori = row[4]
        teks = row[5]

        # Sama dengan script penelitian:
        # hanya baris yang memiliki teks yang diambil.
        if teks:

            rows.append({
                "sumber": "Maret2026",
                "tanggal": tanggal,
                "teks": str(teks).strip(),
                "kategori": (
                    str(kategori).strip()
                    if kategori
                    else "Tidak Diketahui"
                ),
            })

    workbook.close()

    return rows


# ==========================================================
# MEMBACA DATASET APRIL
# ==========================================================

def read_april(file_path):
    """
    Membaca REKAP_FO_LOKET_BULAN_April_2026.xlsx.

    Struktur:
    - Sheet    : Sheet7
    - Tanggal  : kolom A
    - Kategori : kolom F
    - Teks     : kolom G
    - Mulai    : baris 7
    """

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    sheet_name = "Sheet7"

    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' tidak ditemukan "
            f"pada dataset April."
        )

    worksheet = workbook[
        sheet_name
    ]

    rows = []

    for row in worksheet.iter_rows(
        min_row=7,
        values_only=True,
    ):

        tanggal = row[0]
        kategori = row[5]
        teks = row[6]

        if teks:

            rows.append({
                "sumber": "April2026",
                "tanggal": tanggal,
                "teks": str(teks).strip(),
                "kategori": (
                    str(kategori).strip()
                    if kategori
                    else "Tidak Diketahui"
                ),
            })

    workbook.close()

    return rows


# ==========================================================
# MEMBACA DATASET MEI
# ==========================================================

def read_mei(file_path):
    """
    Membaca REKAP_FO_LOKET_BULAN_MEI_2026.xlsx.

    Struktur:
    - Sheet    : MEI 2026 - LAPORAN DATA PELAYAN
    - Tanggal  : kolom A
    - Kategori : kolom F
    - Teks     : kolom G
    - Mulai    : baris 7
    """

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    sheet_name = (
        "MEI 2026 - LAPORAN DATA PELAYAN"
    )

    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' tidak ditemukan "
            f"pada dataset Mei."
        )

    worksheet = workbook[
        sheet_name
    ]

    rows = []

    for row in worksheet.iter_rows(
        min_row=7,
        values_only=True,
    ):

        tanggal = row[0]
        kategori = row[5]
        teks = row[6]

        if teks:

            rows.append({
                "sumber": "Mei2026",
                "tanggal": tanggal,
                "teks": str(teks).strip(),
                "kategori": (
                    str(kategori).strip()
                    if kategori
                    else "Tidak Diketahui"
                ),
            })

    workbook.close()

    return rows


# ==========================================================
# MEMBACA DATASET PEMAKAMAN
# ==========================================================

def read_pemakaman(file_path):
    """
    Membaca dataset Pelayanan Informasi Pemakaman.

    Struktur mengikuti 01_preprocessing.py penelitian:
    - Sheet         : Form Responses 1
    - Timestamp     : kolom A
    - Jenis Layanan : kolom H
    - Teks          : kolom I
    - Kategori      : kolom M
    - Mulai         : baris 2

    Hanya data dengan Jenis Layanan = Informasi.
    """

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    sheet_name = "Form Responses 1"

    if sheet_name not in workbook.sheetnames:
        workbook.close()

        raise ValueError(
            f"Sheet '{sheet_name}' tidak ditemukan "
            f"pada dataset Pemakaman."
        )

    worksheet = workbook[sheet_name]

    rows = []

    # max_col=13 memastikan tuple selalu mencakup
    # kolom A sampai M.
    for row in worksheet.iter_rows(
        min_row=2,
        min_col=1,
        max_col=13,
        values_only=True,
    ):

        tanggal = row[0]      # A
        jenis = row[7]        # H
        teks = row[8]         # I
        kategori = row[12]    # M

        # Sama dengan script penelitian:
        # hanya layanan Informasi yang digunakan.
        if (
            teks
            and jenis
            and str(jenis).strip() == "Informasi"
        ):

            rows.append({
                "sumber": "Pemakaman",
                "tanggal": tanggal,
                "teks": str(teks).strip(),
                "kategori": (
                    str(kategori).strip()
                    if kategori
                    else "Pemakaman Umum"
                ),
            })

    workbook.close()

    return rows

# ==========================================================
# MEMBACA SATU DATASET
# ==========================================================

def load_standardized_dataset(file_path):
    """
    Membaca satu dataset berdasarkan jenis sumbernya
    menggunakan struktur yang sama dengan
    01_preprocessing.py penelitian.
    """

    file_path = validate_excel_file(
        file_path
    )

    sumber = determine_source(
        file_path
    )

    if sumber == "Maret2026":
        rows = read_maret(
            file_path
        )

    elif sumber == "April2026":
        rows = read_april(
            file_path
        )

    elif sumber == "Mei2026":
        rows = read_mei(
            file_path
        )

    elif sumber == "Pemakaman":
        rows = read_pemakaman(
            file_path
        )

    else:
        raise ValueError(
            f"Sumber dataset tidak didukung: "
            f"{sumber}"
        )

    return pd.DataFrame(
        rows,
        columns=[
            "sumber",
            "tanggal",
            "teks",
            "kategori",
        ],
    )


# ==========================================================
# MENGGABUNGKAN DATASET
# ==========================================================

def combine_datasets(file_paths):
    """
    Membaca dan menggabungkan seluruh dataset historis.

    Mekanisme mengikuti 01_preprocessing.py penelitian.
    """

    if not file_paths:
        raise ValueError(
            "Tidak ada dataset yang diberikan."
        )

    datasets = []

    for file_path in file_paths:

        dataset = load_standardized_dataset(
            file_path
        )

        if dataset.empty:
            raise ValueError(
                f"Dataset tidak menghasilkan data: "
                f"{Path(file_path).name}"
            )

        datasets.append(
            dataset
        )

    combined = pd.concat(
        datasets,
        ignore_index=True,
    )

    if combined.empty:
        raise ValueError(
            "Dataset gabungan kosong."
        )

    return combined